from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from depipeline.errors import UnsupportedFileTypeError
from depipeline.extraction import extract_text_with_diagnostics

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None


class ExtractionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(__file__).resolve().parents[1]

    def test_core_sample_fixtures_exist(self) -> None:
        core_dir = self.root / "fixtures" / "core"
        expected = ["sample.txt", "sample.docx", "sample.pdf"]

        missing = [name for name in expected if not (core_dir / name).exists()]
        self.assertEqual([], missing, f"Missing core fixtures: {missing}")

    @unittest.skipIf(Document is None, "python-docx not installed")
    def test_docx_file_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            file_path = Path(tmp_dir) / "sample.docx"

            doc = Document()
            doc.add_paragraph("Participant: Casey")
            doc.add_paragraph("Suggested improvement: show progress steps")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "sentiment"
            table.cell(0, 1).text = "negative"
            doc.save(file_path)

            diagnostics = extract_text_with_diagnostics(file_path)

            self.assertIn("Participant: Casey", diagnostics.text)
            self.assertIn("Suggested improvement: show progress steps", diagnostics.text)
            self.assertIn("sentiment", diagnostics.text)
            self.assertFalse(diagnostics.is_scanned_pdf)
            self.assertEqual([], diagnostics.warnings)

    @unittest.skipIf(Document is None, "python-docx not installed")
    def test_fixture_sample_docx_is_extracted(self) -> None:
        file_path = self.root / "fixtures" / "core" / "sample.docx"
        self.assertTrue(file_path.exists(), "Expected fixture sample.docx to exist")

        diagnostics = extract_text_with_diagnostics(file_path)

        self.assertIn("Participant: Alex", diagnostics.text)
        self.assertIn("Suggested improvement", diagnostics.text)
        self.assertFalse(diagnostics.is_scanned_pdf)
        self.assertEqual([], diagnostics.warnings)

    def test_markdown_file_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            file_path = Path(tmp_dir) / "notes.md"
            file_path.write_text(
                "---\n" "title: Notes\n" "---\n\n" "# Header\n\n" "- export CSV fails\n",
                encoding="utf-8",
            )

            diagnostics = extract_text_with_diagnostics(file_path)

            self.assertIn("Header", diagnostics.text)
            self.assertIn("export CSV fails", diagnostics.text)
            self.assertNotIn("title: Notes", diagnostics.text)
            self.assertFalse(diagnostics.is_scanned_pdf)
            self.assertEqual([], diagnostics.warnings)

    def test_unsupported_extension_still_raises(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            file_path = Path(tmp_dir) / "notes.rtf"
            file_path.write_text("hello", encoding="utf-8")

            with self.assertRaises(UnsupportedFileTypeError):
                extract_text_with_diagnostics(file_path)


if __name__ == "__main__":
    unittest.main()
