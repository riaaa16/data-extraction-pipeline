from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import List, Union

from .errors import ExtractionError, ScannedPDFError, UnsupportedFileTypeError
from .normalize import normalize_text

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

try:
    from docx import Document  # python-docx
except Exception:  # pragma: no cover
    Document = None


PathLike = Union[str, Path]


@dataclass(frozen=True)
class ExtractionDiagnostics:
    text: str
    warnings: List[str]
    is_scanned_pdf: bool


def extract_text(path: PathLike) -> str:
    """Extract readable text from a supported file.

    Supported:
    - .txt
    - .docx
    - .pdf (text-based; scanned PDFs are detected and rejected)

    Returns the extracted text as a string.
    """

    diagnostics = extract_text_with_diagnostics(path)
    if diagnostics.is_scanned_pdf:
        raise ScannedPDFError(
            "PDF appears to be scanned/image-based (no extractable text). "
            "OCR is not supported in Sprint 01."
        )
    return diagnostics.text


def extract_text_with_diagnostics(path: PathLike) -> ExtractionDiagnostics:
    file_path = Path(path)
    if not file_path.exists():
        raise ExtractionError(f"File not found: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        text = _extract_txt(file_path)
        return ExtractionDiagnostics(text=text, warnings=[], is_scanned_pdf=False)

    if suffix == ".docx":
        text = _extract_docx(file_path)
        return ExtractionDiagnostics(text=text, warnings=[], is_scanned_pdf=False)

    if suffix == ".pdf":
        return _extract_pdf(file_path)

    raise UnsupportedFileTypeError(f"Unsupported file type: {suffix} ({file_path.name})")


def _extract_txt(path: Path) -> str:
    encodings_to_try = ["utf-8", "utf-8-sig", "cp1252", "latin-1"]
    last_error: Exception | None = None

    for encoding in encodings_to_try:
        try:
            text = path.read_text(encoding=encoding)
            return normalize_text(text)
        except Exception as exc:  # pragma: no cover
            last_error = exc

    try:
        with path.open("r", encoding="utf-8", errors="replace") as f:
            return normalize_text(f.read())
    except Exception as exc:
        raise ExtractionError(f"Failed reading text file: {path}") from (last_error or exc)


def _extract_docx(path: Path) -> str:
    if Document is None:
        raise ExtractionError(
            "python-docx is not installed. Install dependencies with: pip install -r requirements.txt"
        )

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Failed opening DOCX: {path}") from exc

    blocks: List[str] = []

    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text or "")
        if text:
            blocks.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = normalize_text(cell.text or "")
                if text:
                    blocks.append(text)

    return normalize_text("\n\n".join(blocks))


def _extract_pdf(path: Path) -> ExtractionDiagnostics:
    if fitz is None:
        raise ExtractionError(
            "PyMuPDF is not installed. Install dependencies with: pip install -r requirements.txt"
        )

    warnings: List[str] = []

    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ExtractionError(f"Failed opening PDF: {path}") from exc

    try:
        page_texts: List[str] = []
        pages_with_text = 0

        for page in doc:
            page_text = page.get_text("text") or ""
            page_text = normalize_text(page_text)

            # Consider a page to have text if it contains at least a few non-whitespace chars.
            non_ws = len("".join(page_text.split()))
            if non_ws >= 20:
                pages_with_text += 1

            if page_text:
                page_texts.append(page_text)

        text = normalize_text("\n\n".join(page_texts))
        if pages_with_text == 0:
            warnings.append(
                "No extractable text detected. PDF may be scanned/image-based (OCR not supported yet)."
            )
            return ExtractionDiagnostics(text="", warnings=warnings, is_scanned_pdf=True)

        if pages_with_text < max(1, doc.page_count // 2):
            warnings.append(
                "Some pages had little/no extractable text; output may be incomplete (possible scanned pages)."
            )

        return ExtractionDiagnostics(text=text, warnings=warnings, is_scanned_pdf=False)
    finally:
        doc.close()
